"""Write the Practitioner Information hand-off DOCX from a small JSON file.

Usage: python3 write_practitioner_info.py <info.json> <output_folder> [final_report.docx]
If the final report is given, every placeholder left in it ([...] brackets, "XX", "Figure XX", "Table XX") is listed automatically (Python, not Claude).
See prompts/practitioner_information.md for the JSON shape.
"""
import json, os, sys
from docx import Document

HEADINGS = [("check", "1. Check / verify"), ("uncertain", "2. Uncertain"),
            ("not_found", "3. Not found / not available"), ("incomplete", "4. Incomplete")]

def scan_placeholders(docx_path):
    import re
    found = []
    d = Document(docx_path)
    paras = [p.text for p in d.paragraphs]
    for t in (c.text for tb in d.tables for r in tb.rows for c in r.cells):
        paras.append(t)
    pat = re.compile(r"\[[A-Z][^\]]{3,}\]|\bXX+\b|Figure XX|Tables? XX|\bXXX\b")
    for t in paras:
        hits = pat.findall(t)
        if hits:
            found.append((sorted(set(hits)), " ".join(t.split())[:90]))
    return found


def main(info_path, out_dir, report_docx=None):
    info = json.load(open(info_path, encoding="utf-8"))
    doc = Document()
    doc.add_heading("Practitioner Information", level=1)
    doc.add_paragraph(f"Site: {info.get('site', '')}    Report: {info.get('report', '')}")
    doc.add_paragraph("Hand-off notes for the practitioner. Not part of the report; do not issue with it.")
    for key, title in HEADINGS:
        doc.add_heading(title, level=2)
        items = info.get(key) or []
        if not items:
            doc.add_paragraph("None identified.")
        for it in items:
            doc.add_paragraph(it, style="List Bullet")
    if report_docx:
        doc.add_heading("5. Placeholders left in the report (auto-scanned)", level=2)
        hits = scan_placeholders(report_docx)
        if not hits:
            doc.add_paragraph("None found.")
        for marks, where in hits:
            doc.add_paragraph(f"{', '.join(m[:60] for m in marks)}  —  \"{where}\"", style="List Bullet")
    site = "".join(c if c.isalnum() or c in "-_" else "_" for c in info.get("site", "site"))
    out = os.path.join(out_dir, f"Practitioner_Information_{site}.docx")
    doc.save(out)
    print("Saved:", out)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
