"""Provider-agnostic LLM client — OpenRouter via the openai SDK.

Config (env vars, or a ``backend/.env`` file — see ``.env.example``):
- ``OPENROUTER_API_KEY`` — required for real calls
- ``REPORT_MODEL`` — default model (e.g. ``anthropic/claude-sonnet-4-*``)
- ``REPORT_MODEL_EXTRACT`` — model for the site-facts extraction stage
  (falls back to ``REPORT_MODEL``)
- ``REPORT_MODEL_AUTHOR`` — model for section authoring (falls back to
  ``REPORT_MODEL``)

Per-stage model switching is a first-class feature: OpenRouter is a single
endpoint and the model is a per-request parameter, so each pipeline stage can
target a different model (e.g. a vision-capable Claude model for PDF
extraction, a cheaper DeepSeek model for the many authoring calls).
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Iterable

from dotenv import load_dotenv
from openai import OpenAI

# Load backend/.env if present (gitignored) so the key is set once, not per
# terminal. Real env vars always win over the .env file.
_BACKEND_ROOT = Path(__file__).resolve().parents[2]
load_dotenv(_BACKEND_ROOT / ".env")

OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"

# Default model: a Claude model via OpenRouter to mimic the current Claude
# Desktop setup (see docs/ai-reporting-plan.md decision 8).
DEFAULT_MODEL = "anthropic/claude-sonnet-5"


class LLMError(RuntimeError):
    """Raised when an LLM call fails or no API key is configured."""


def _env(name: str, default: str | None = None) -> str | None:
    return os.environ.get(name) or default


def model_for(stage: str) -> str:
    """Resolve the model for a pipeline stage (extract/author/other)."""
    if stage == "extract":
        return _env("REPORT_MODEL_EXTRACT") or _env("REPORT_MODEL") or DEFAULT_MODEL
    if stage == "author":
        return _env("REPORT_MODEL_AUTHOR") or _env("REPORT_MODEL") or DEFAULT_MODEL
    return _env("REPORT_MODEL") or DEFAULT_MODEL


class LLMClient:
    """Thin wrapper around the OpenRouter chat-completions API."""

    def __init__(self, api_key: str | None = None) -> None:
        self.api_key = api_key or _env("OPENROUTER_API_KEY")
        self._client: OpenAI | None = None

    @property
    def client(self) -> OpenAI:
        if self._client is None:
            if not self.api_key:
                raise LLMError(
                    "OPENROUTER_API_KEY is not set. Set it in the environment "
                    "to enable report generation."
                )
            self._client = OpenAI(
                api_key=self.api_key, base_url=OPENROUTER_BASE_URL
            )
        return self._client

    # -- chat completions -------------------------------------------------
    def complete(
        self,
        system: str,
        user: str,
        *,
        model: str | None = None,
        temperature: float = 0.2,
        max_tokens: int = 8192,
    ) -> str:
        """Single chat completion; returns the assistant text."""
        try:
            resp = self.client.chat.completions.create(
                model=model or model_for("other"),
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": user},
                ],
                temperature=temperature,
                max_tokens=max_tokens,
            )
        except Exception as exc:  # openai raises a rich error hierarchy
            raise LLMError(f"LLM call failed: {exc}") from exc
        content = resp.choices[0].message.content
        if content is None:
            raise LLMError("LLM returned an empty response.")
        return content

    def complete_json(
        self,
        system: str,
        user: str,
        *,
        model: str | None = None,
        temperature: float = 0.0,
        max_tokens: int = 8192,
        retries: int = 3,
    ) -> dict[str, Any]:
        """Chat completion with JSON output; parses and returns the object.

        Uses ``response_format={"type": "json_object"}`` on the first attempt
        (falls back to plain completion if the model rejects it), and retries
        on empty responses, API errors, and unparseable JSON — models
        occasionally return an empty string, which is not a real answer.
        """
        import json
        import re

        kwargs: dict[str, Any] = dict(
            model=model or model_for("other"),
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            temperature=temperature,
            max_tokens=max_tokens,
        )
        last_error: Exception | None = None
        for attempt in range(retries):
            try:
                if attempt == 0:
                    kwargs["response_format"] = {"type": "json_object"}
                else:
                    kwargs.pop("response_format", None)
                resp = self.client.chat.completions.create(**kwargs)
            except Exception as exc:
                last_error = exc
                continue
            content = resp.choices[0].message.content or ""
            if not content.strip():
                last_error = LLMError("LLM returned an empty response.")
                continue
            try:
                return json.loads(content)
            except json.JSONDecodeError:
                # Strip markdown fences if the model wrapped the JSON (object
                # or array).
                match = re.search(r"```(?:json)?\s*([\[{].*?[\]}])\s*```", content, re.DOTALL)
                if match:
                    return json.loads(match.group(1))
                last_error = LLMError(f"LLM did not return valid JSON: {content[:500]}")
                continue
        raise LLMError(f"LLM call failed after {retries} attempts: {last_error}")


# -- document ingestion -----------------------------------------------------


def ingest_documents(paths: Iterable[Path], max_chars_per_doc: int = 60_000) -> list[dict[str, Any]]:
    """Read PDF/Word/text documents into a list of {name, text} blocks.

    PDFs are read with pypdf; .docx with python-docx; plain text read as-is.
    Each document's text is truncated to ``max_chars_per_doc`` to bound the
    prompt size (the extraction call is a single LLM call).
    """
    blocks: list[dict[str, Any]] = []
    for path in paths:
        path = Path(path)
        if not path.exists():
            continue
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                text = _read_pdf(path)
            elif suffix in (".docx", ".doc"):
                text = _read_docx(path)
            else:
                text = path.read_text(encoding="utf-8", errors="replace")
        except Exception as exc:  # keep going; flag the doc as unreadable
            blocks.append({"name": path.name, "text": "", "error": str(exc)})
            continue
        blocks.append({"name": path.name, "text": text[:max_chars_per_doc]})
    return blocks


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def assemble_document_payload(blocks: list[dict[str, Any]]) -> str:
    """Format ingested document blocks into a single user-message payload."""
    sections = []
    for block in blocks:
        header = f"### {block['name']}"
        if block.get("error"):
            sections.append(f"{header}\n[unreadable: {block['error']}]")
        elif block["text"]:
            sections.append(f"{header}\n{block['text']}")
        else:
            sections.append(f"{header}\n[no extractable text]")
    return "\n\n".join(sections)