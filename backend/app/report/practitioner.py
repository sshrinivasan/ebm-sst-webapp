"""Step 6 — practitioner information hand-off DOCX.

Builds the four-category hand-off (check / uncertain / not_found /
incomplete) from the session's extraction flags, intake answers, and a scan
of the final DOCX for leftover placeholders, then runs
``write_practitioner_info.py`` to produce ``Practitioner_Information_<site>.docx``.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path
from typing import Any

from .session import ReportSession

_SKILL_ROOT = Path(__file__).resolve().parents[2] / "report_skill"
WRITER_SCRIPT = _SKILL_ROOT / "scripts" / "write_practitioner_info.py"

# Placeholder patterns left in the final DOCX that the practitioner must fill.
_PLACEHOLDER_PATTERNS = [
    r"XX",
    r"Figure XX",
    r"Tables? XX",
    r"\[SITE-CONTEXT PLACEHOLDER\]",
    r"\[PRACTITIONER",
    r"<INSERT_[A-Z0-9_]+>",
]


class PractitionerError(RuntimeError):
    """Raised when the practitioner hand-off cannot be built."""


def _scan_placeholders(docx_path: Path) -> list[str]:
    """Find leftover placeholder strings in the final DOCX."""
    from docx import Document

    found: set[str] = set()
    try:
        doc = Document(str(docx_path))
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    texts.append(cell.text)
        for text in texts:
            for pattern in _PLACEHOLDER_PATTERNS:
                if re.search(pattern, text):
                    found.add(pattern)
    except Exception:
        pass
    return sorted(found)


def build_practitioner_info(session: ReportSession) -> dict[str, Any]:
    """Assemble the practitioner_info.json payload from session state."""
    facts = session.site_facts or {}
    site = facts.get("site_lsd") or facts.get("site_name") or "site"
    report_name = Path(session.final_docx or "report.docx").name

    check: list[str] = []
    uncertain: list[str] = []
    not_found: list[str] = []
    incomplete: list[str] = []

    # Extraction flags -> not found / uncertain.
    for flag in session.extraction_flags:
        message = str(flag.get("message", ""))
        field = flag.get("field")
        if not message:
            continue
        if flag.get("ambiguous"):
            uncertain.append(f"Extraction: {message}" + (f" (field: {field})" if field else ""))
        else:
            not_found.append(f"Extraction: {message}" + (f" (field: {field})" if field else ""))

    # Intake Q18 (Claude-authored judgment) -> check.
    q18 = session.intake_answers.get("q18")
    if q18:
        check.append(
            "Section 5.2 closing transport interpretation was drafted from "
            f"intake Q18: \"{q18}\" — confirm the interpretation."
        )

    # Defaulted/blank intake answers -> uncertain.
    for qid, value in session.intake_answers.items():
        if value is None or (isinstance(value, str) and not value.strip()):
            uncertain.append(f"Intake {qid.upper()} was left blank; a default was used.")

    # Leftover placeholders in the final DOCX -> incomplete.
    if session.final_docx:
        for pattern in _scan_placeholders(Path(session.final_docx)):
            incomplete.append(f"Placeholder '{pattern}' remains in the report.")

    return {
        "site": str(site),
        "report": report_name,
        "check": check,
        "uncertain": uncertain,
        "not_found": not_found,
        "incomplete": incomplete,
    }


def write_practitioner_docx(
    session: ReportSession,
    *,
    timeout: int = 120,
) -> Path:
    """Build and save the practitioner hand-off DOCX (Step 6)."""
    if not WRITER_SCRIPT.exists():
        raise PractitionerError(f"write_practitioner_info.py not found: {WRITER_SCRIPT}")

    info = build_practitioner_info(session)
    session.practitioner_info = info
    info_path = session.save_practitioner_info()

    cmd = [
        sys.executable,
        str(WRITER_SCRIPT),
        str(info_path),
        str(session.dir),
    ]
    try:
        proc = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=timeout,
            cwd=str(_SKILL_ROOT),
        )
    except subprocess.TimeoutExpired as exc:
        raise PractitionerError(f"Practitioner writer timed out after {timeout}s.") from exc

    if proc.returncode != 0:
        raise PractitionerError(
            f"Practitioner writer failed (exit {proc.returncode}):\n{proc.stderr[-2000:]}"
        )

    # Locate the produced DOCX (named Practitioner_Information_<site>.docx).
    matches = sorted(session.dir.glob("Practitioner_Information_*.docx"))
    if not matches:
        raise PractitionerError("Practitioner writer did not produce a DOCX.")
    session.practitioner_docx = str(matches[-1])
    return matches[-1]